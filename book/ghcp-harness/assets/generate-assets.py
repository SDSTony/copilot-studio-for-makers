from pathlib import Path
import subprocess
import sys
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parent
HEADERS = [
    "PlanId", "Destination", "StartDate", "EstimatedTotal", "ReviewResult", "Status",
]
ROWS = [
    ["SAMPLE-001", "광주", "2026-12-03", 287000, "한도 초과", "준비 중"],
    ["SAMPLE-002", "전주", "2026-12-08", 73000, "한도 이내", "완료"],
]


def generate_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TripPlans"
    sheet.append(HEADERS)
    for row in ROWS:
        sheet.append(row)
    table = Table(displayName="TripPlans", ref="A1:F3")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 22
        for cell in column:
            cell.font = Font(name="Arial", size=11, bold=cell.row == 1)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row > 1:
                cell.comment = Comment(
                    "Source: GHCP 출장 실습 가상 샘플, LAB-2026-01. 실제 출장 기록이 아닙니다.",
                    "Copilot Studio for Makers",
                )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column == 3:
                cell.number_format = "@"
            elif cell.column == 4:
                cell.number_format = '#,##0" 원"'
    notes = workbook.create_sheet("ReadMe")
    notes.append(["교육용 출장 계획 기록"])
    notes.append(["가상 데이터입니다. 실제 개인정보나 출장 기록을 넣지 마세요."])
    notes.append(["EstimatedTotal은 검토 당시 입력받은 총액의 기록값이며 실시간 계산 모델이 아닙니다."])
    notes.append(["6개 열: PlanId, Destination, StartDate, EstimatedTotal, ReviewResult, Status"])
    notes.append(["날짜는 YYYY-MM-DD 텍스트, 금액은 KRW 숫자, 테이블 이름은 TripPlans입니다."])
    notes.append(["한도 이내/초과는 숙박 검토 결과이며 승인 상태가 아닙니다."])
    notes.append(["기타 비용에 한도 기준은 없으며, 용도와 적정성은 별도 확인합니다."])
    notes.append(["연결 중 동시 편집을 피하고 조회/등록 후 반영을 기다리세요."])
    notes.column_dimensions["A"].width = 100
    for row in notes:
        row[0].font = Font(name="Arial", size=11)
        row[0].alignment = Alignment(wrap_text=True)
    workbook.save(ROOT / "trip-plans.xlsx")


def generate_skill():
    source = ROOT / "trip-plan-report"
    normalize_presentation(source / "assets" / "trip-plan-template.pptx")
    with ZipFile(ROOT / "trip-plan-report.zip", "w", ZIP_DEFLATED) as package:
        for file in sorted(source.rglob("*")):
            if file.is_file():
                package.write(file, file.relative_to(source).as_posix())


def normalize_presentation(path):
    with ZipFile(path) as source:
        document = ET.fromstring(source.read("ppt/presentation.xml"))
        ordering = [child.tag.split("}")[-1] for child in document]
    if ordering.index("notesMasterIdLst") < ordering.index("sldIdLst"):
        return
    # A root-only XML reorder passes XSD but breaks PowerPoint; use a native package roundtrip.
    if sys.platform != "win32":
        raise RuntimeError("Regenerating the PPTX requires Windows and desktop PowerPoint.")
    subprocess.run([
        "powershell.exe", "-NoProfile", "-ExecutionPolicy", "RemoteSigned",
        "-File", str(ROOT / "normalize-presentation.ps1"),
        "-Path", str(path),
    ], check=True)


if __name__ == "__main__":
    generate_workbook()
    generate_skill()
    print("Generated trip-plans.xlsx and trip-plan-report.zip")
